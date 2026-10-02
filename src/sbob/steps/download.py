"""Passo download: registrazioni Polimi → video/<stem>.mp4.

Usa il downloader `prd` nel SUO venv (settings.downloader/.venv): non si importa, perché vincola typer 0.6.

Flusso in due fasi, così una ripetizione non riscarica nulla:
  1. piano   `prd <tipo> ... --no-aria2c` → xlsx con data e link (nessun download)
  2. scarico `prd txt ids.txt` con SOLO gli id non ancora nel manifest, in staging/dl

Identità di una registrazione = nome che prd le dà, "YYYY-MM-DD HH-MM" (manifest.videos: chiave → stem).
"""

from __future__ import annotations

import os
import re
import shutil
import subprocess
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from sbob.config import Course, Settings
from sbob.core import naming
from sbob.core.batch import atomic_write_text, list_inputs
from sbob.core.report import NeedsHuman, StepReport
from sbob.core.status import VIDEO_EXT
from sbob.steps.base import StepContext

_ID_RE = re.compile(r"[0-9a-f]{32}", re.I)
_TICKET_MSGS = ("refresh",                    # prd: "Try refreshing the ticket"
                "downloadrecordinginfo")      # KeyError di prd quando Webex risponde senza dati (ticket non valido)


def prd_python(settings: Settings) -> Path:
    py = settings.downloader / ".venv" / "bin" / "python"
    if not py.exists():
        raise NeedsHuman(f"Downloader non trovato in {settings.downloader}",
                         action="imposta `downloader` in sbob.toml (clone di polimi_recordings_downloader con .venv)")
    return py


def _run_prd(settings: Settings, args: list[str], cwd: Path) -> subprocess.CompletedProcess:
    cwd.mkdir(parents=True, exist_ok=True)
    # prd non è installato nel suo venv: si importa dalla cartella del clone
    env = {**os.environ, "PYTHONPATH": os.pathsep.join(filter(None, [str(settings.downloader),
                                                                       os.environ.get("PYTHONPATH")]))}
    return subprocess.run([str(prd_python(settings)), "-m", "prd", *args], cwd=cwd, text=True,
                          capture_output=True, env=env)


def set_cookie(settings: Settings, nome: str, valore: str) -> None:
    res = _run_prd(settings, ["set-cookie", nome, valore], settings.downloader)
    if res.returncode != 0:
        raise RuntimeError((res.stdout + res.stderr).strip() or "set-cookie fallito")


def read_links_file(path: Path) -> tuple[list[str], dict[str, dict]]:
    """File di link, anche arricchito: `link<TAB>dd/mm/yyyy HH:MM<TAB>forma didattica<TAB>argomento`.
    Righe vuote e `#` sono ignorate. Restituisce (link puliti per prd, {video_id: {data, tipo, argomento}})."""
    links, meta = [], {}
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip() or line.lstrip().startswith("#"):
            continue
        cols = [c.strip() for c in line.split("\t")]
        links.append(cols[0])
        if len(cols) > 1 and (m := _ID_RE.search(cols[0])):
            info: dict = {}
            try:
                info["data"] = datetime.strptime(cols[1], "%d/%m/%Y %H:%M")
            except ValueError:
                pass
            if len(cols) > 2 and cols[2]:
                info["tipo"] = FORMA_TO_TIPO.get(cols[2].lower())
            if len(cols) > 3 and cols[3]:
                info["argomento"] = cols[3]
            meta[m.group(0).lower()] = info
    return links, meta


def _source_args(course: Course, options: dict, course_dir: Path, staging: Path | None = None) -> tuple[list[str], dict]:
    src = course.sorgente
    tipo = src.get("tipo")
    if not tipo:
        raise NeedsHuman(f"Il corso '{course.slug}' non ha una sorgente", action=f"aggiungi sorgente in [corsi.{course.slug}]")
    common = ["--course", course.slug, "--academic-year", course.anno_accademico]
    if tipo in ("txt", "webpage-html"):
        file = Path(options.get("links") or src.get("file", "link.txt")).expanduser()
        file = file if file.is_absolute() else course_dir / file
        if not file.is_file():
            raise NeedsHuman(f"File link non trovato: {file}", action=f"crea {file} con un link Webex per riga")
        if tipo == "txt" and staging is not None:
            links, meta = read_links_file(file)      # a prd vanno solo i link; i metadati li usa sbob
            clean = staging / "links.txt"
            atomic_write_text(clean, "\n".join(links) + "\n")
            return [tipo, str(clean), *common], meta
        return [tipo, str(file), *common], {}
    if tipo == "webpage-url":
        return [tipo, src["url"], *common], {}
    return [tipo, src["url"]], {}  # webeep, archives: corso e anno li ricava prd dalla pagina


# Quali cookie servono a prd per ogni sorgente: ticket (Webex) sempre, più quello del sito da cui si leggono i link
COOKIES_FOR_SOURCE = {
    "webeep": ("ticket", "MoodleSession"),
    "archives": ("ticket", "JSESSIONID", "INGRESSCOOKIE"),   # archivio su onlineservices.polimi.it
}
_COOKIE_WHERE = {"ticket": "politecnicomilano.webex.com", "MoodleSession": "webeep.polimi.it",
                 "JSESSIONID": "onlineservices.polimi.it (archivio registrazioni)",
                 "INGRESSCOOKIE": "onlineservices.polimi.it (archivio registrazioni)",
                 "SSL_JSESSIONID": "www11.ceda.polimi.it (vecchio archivio)"}
# "Forma didattica" dell'archivio recman (prd la mette nel subject come "[Forma] argomento") → tipo di lezione
FORMA_TO_TIPO = {"lezione": "lez", "esercitazione": "ese", "laboratorio": "lab", "altro": "sem"}
_FORMA = re.compile(r"^\[([^\]]+)\]\s*(.*)$", re.S)


def cookie_help(tipo: str | None) -> str:
    needed = COOKIES_FOR_SOURCE.get(tipo or "", ("ticket",))
    return "; ".join(f"copia il cookie '{c}' da {_COOKIE_WHERE[c]} e lancia: sbob cookie {c} <valore>" for c in needed)


def _check_output(res: subprocess.CompletedProcess, tipo: str | None = None) -> None:
    text = (res.stdout or "") + (res.stderr or "")
    lowered = text.lower()
    if res.returncode != 0 and ("moodlesession" in lowered or "jsessionid" in lowered or "cookie" in lowered
                                or "zero recordings" in lowered or any(m in lowered for m in _TICKET_MSGS)):
        needed = " e ".join(COOKIES_FOR_SOURCE.get(tipo or "", ("ticket",)))
        raise NeedsHuman(f"Cookie scaduti o mancanti (per questa sorgente servono: {needed})",
                         action=cookie_help(tipo))
    if res.returncode != 0:
        tail = [l for l in text.strip().splitlines() if l.strip()][-3:]
        raise RuntimeError("prd fallito: " + " | ".join(tail))


def read_plan(staging_plan: Path) -> dict[str, dict[str, str | None]]:
    """{ 'YYYY-MM-DD HH-MM': {"id", "tipo", "argomento"} } da tutti gli xlsx generati da prd.
    `tipo` è dedotto dalla forma didattica quando la sorgente la fornisce (archivio recman), altrimenti None."""
    found: dict[str, dict[str, str | None]] = {}
    for xlsx in sorted(staging_plan.rglob("*.xlsx")):
        ws = load_workbook(xlsx).active
        for row in ws.iter_rows(min_row=2):
            link = row[0].hyperlink.target if row[0].hyperlink else None
            when = row[2].value
            if not link or not when or not (m := _ID_RE.search(link)):
                continue
            dt = when if isinstance(when, datetime) else datetime.strptime(str(when), "%Y-%m-%d %H:%M")
            subject = str(row[3].value or "").strip() if len(row) > 3 else ""
            tipo, argomento = None, subject or None
            if fm := _FORMA.match(subject):
                tipo = FORMA_TO_TIPO.get(fm.group(1).strip().lower())
                argomento = fm.group(2).strip() or None
            found[dt.strftime("%Y-%m-%d %H-%M")] = {"id": m.group(0), "tipo": tipo, "argomento": argomento}
    return found


def assign_names(course: Course, existing_stems: list[str], new: dict[str, str],
                 dates: dict[str, datetime] | None = None) -> dict[str, str]:
    """{key prd: tipo} → {key prd: stem canonico}, numerando per tipo dopo l'ultima lezione esistente di quel tipo.
    `dates` sovrascrive la data dedotta dal nome prd (che viene da Webex e a volte è sbagliata)."""
    parsed = [n for s in existing_stems if (n := naming.parse(s))]
    out: dict[str, str] = {}
    for tipo in sorted(set(new.values())):
        keys = [k for k, t in new.items() if t == tipo]
        dated = sorted(((dates or {}).get(k) or naming.parse_prd(k), k) for k in keys)
        names = naming.next_numbers(parsed, [d.date() for d, _ in dated], course.slug, tipo)
        out.update({k: n.stem for (_, k), n in zip(dated, names)})
    return out


def run(ctx: StepContext) -> StepReport:
    rep = ctx.report("download")
    lay, course = ctx.layout, ctx.course
    tipo = ctx.options.get("tipo")          # None = dalla forma didattica se c'è, altrimenti "lez"
    if tipo is not None and tipo not in naming.TIPI:
        rep.error = f"tipo '{tipo}' non valido ({', '.join(naming.TIPI)})"
        return rep
    manifest = ctx.manifest()

    # 1. piano
    plan_dir = lay.staging / "plan"
    shutil.rmtree(plan_dir, ignore_errors=True)
    args, link_meta = _source_args(course, ctx.options, course.cartella, lay.staging)
    ctx.log("download: leggo l'elenco delle registrazioni…")
    _check_output(_run_prd(ctx.settings, [*args, "--no-aria2c", "--create-xlsx", "--output", str(plan_dir)],
                           lay.state), course.sorgente.get("tipo"))
    plan = read_plan(plan_dir)
    for info in plan.values():                   # metadati del file di link (archivio): prevalgono
        extra = link_meta.get(info["id"].lower(), {})
        info["tipo"] = extra.get("tipo") or info["tipo"]
        info["argomento"] = extra.get("argomento") or info["argomento"]
        info["data"] = extra.get("data")
    if not plan:
        rep.notes.append("Nessuna registrazione trovata nella sorgente.")
        return rep

    known = set() if ctx.force else set(manifest.videos)
    new_keys = sorted(k for k in plan if k not in known)
    rep.skipped = sorted(manifest.videos[k] for k in plan if k in known)
    existing = [p.stem for p in list_inputs(lay.video, VIDEO_EXT)]
    names = assign_names(course, existing, {k: tipo or plan[k]["tipo"] or "lez" for k in new_keys},
                         {k: plan[k]["data"] for k in new_keys if plan[k].get("data")})
    if ctx.only:
        names = {k: s for k, s in names.items() if s in ctx.only}
    if ctx.dry_run or not names:
        rep.done = sorted(names.values())
        return rep

    # 2. scarico dei soli nuovi
    dl_dir = lay.staging / "dl"
    shutil.rmtree(dl_dir, ignore_errors=True)
    dl_dir.mkdir(parents=True)          # prd crea --output solo quando genera l'xlsx (qui disattivato)
    ids = lay.staging / "ids.txt"
    atomic_write_text(ids, "\n".join(plan[k]["id"] for k in names) + "\n")
    ctx.log(f"download: scarico {len(names)} registrazioni…")
    _check_output(_run_prd(ctx.settings, ["txt", str(ids), "--course", course.slug, "--academic-year",
                                          course.anno_accademico, "--no-create-xlsx", "--output", str(dl_dir)],
                           lay.state))
    if leftovers := list(dl_dir.rglob("*.aria2")):
        rep.error = f"download incompleto: {len(leftovers)} file .aria2 rimasti in {dl_dir}; rilancia per riprendere"
        return rep

    lay.ensure("video")
    got = {p.stem: p for p in dl_dir.rglob("*.mp4")}
    for key, stem in names.items():
        src = got.get(key)
        if not src:
            rep.fail(stem, f"file {key}.mp4 non scaricato")
            continue
        dst = lay.video / f"{stem}.mp4"
        shutil.move(str(src), dst)
        manifest.videos[key] = stem
        if plan[key].get("argomento"):
            manifest.meta(stem)["argomento"] = plan[key]["argomento"]
        manifest.save()
        rep.done.append(stem)
        rep.outputs.append(str(dst))
    shutil.rmtree(dl_dir, ignore_errors=True)
    return rep
