from datetime import datetime

import pytest
from openpyxl import Workbook

from sbob.core.layout import Layout
from sbob.core.report import NeedsHuman
from sbob.steps import download
from sbob.steps.base import StepContext

ID1, ID2, ID3 = "a" * 32, "b" * 32, "c" * 32
URL = "https://politecnicomilano.webex.com/recordingservice/sites/politecnicomilano/recording/{}"


def _fake_prd(recordings, calls, subjects=None):
    """Imita prd: con --no-aria2c scrive l'xlsx; con `txt` crea gli mp4 degli id richiesti."""
    def fake(settings, args, cwd):
        calls.append(args)
        out = download.Path(args[args.index("--output") + 1])
        if "--no-aria2c" in args:
            wb = Workbook()
            ws = wb.active
            ws.append(["Link", "Academic year", "Recording date", "Subject"])
            for i, (key, vid) in enumerate(recordings.items(), 2):
                ws.cell(i, 1, "Link").hyperlink = URL.format(vid)
                ws.cell(i, 3, datetime.strptime(key, "%Y-%m-%d %H-%M").strftime("%Y-%m-%d %H:%M"))
                ws.cell(i, 4, (subjects or {}).get(key, ""))
            (out / "x 2025-26").mkdir(parents=True)
            wb.save(out / "x 2025-26" / "x.xlsx")
        else:
            wanted = set(open(args[1]).read().split())
            (out / "prova 2025-26").mkdir(parents=True)
            for key, vid in recordings.items():
                if vid in wanted:
                    (out / "prova 2025-26" / f"{key}.mp4").write_text("video")
        return download.subprocess.CompletedProcess(args, 0, "", "")
    return fake


@pytest.fixture
def course_with_links(settings):
    c = settings.corso("prova")
    c.cartella.mkdir(parents=True)
    (c.cartella / "link.txt").write_text("x\n")
    return c


def test_download_numbers_and_is_incremental(settings, course_with_links, monkeypatch):
    c, calls = course_with_links, []
    recs = {"2025-09-18 10-00": ID2, "2025-09-17 09-30": ID1}
    monkeypatch.setattr(download, "_run_prd", _fake_prd(recs, calls))

    rep = download.run(StepContext(settings, c, quiet=True))
    assert rep.done == ["2025-09-17_prova_lez01", "2025-09-18_prova_lez02"]
    lay = Layout.of(c)
    assert (lay.video / "2025-09-17_prova_lez01.mp4").exists()

    # secondo giro: nulla di nuovo, nessuno scarico
    calls.clear()
    rep2 = download.run(StepContext(settings, c, quiet=True))
    assert rep2.done == [] and len(rep2.skipped) == 2
    assert all("--no-aria2c" in a for a in calls)

    # compare una registrazione nuova: numerazione continua, scarica solo quella
    recs["2025-09-24 16-00"] = ID3
    calls.clear()
    rep3 = download.run(StepContext(settings, c, quiet=True))
    assert rep3.done == ["2025-09-24_prova_lez03"]
    assert open(lay.staging / "ids.txt").read().split() == [ID3]


def test_dry_run_does_not_download(settings, course_with_links, monkeypatch):
    calls = []
    monkeypatch.setattr(download, "_run_prd", _fake_prd({"2025-09-17 09-30": ID1}, calls))
    rep = download.run(StepContext(settings, course_with_links, dry_run=True, quiet=True))
    assert rep.done == ["2025-09-17_prova_lez01"] and rep.dry_run
    assert len(calls) == 1 and not Layout.of(course_with_links).video.exists()


def test_expired_ticket_needs_human():
    res = download.subprocess.CompletedProcess([], 1, "Try refreshing the ticket", "")
    with pytest.raises(NeedsHuman) as e:
        download._check_output(res)
    assert "sbob cookie ticket" in e.value.action
    res = download.subprocess.CompletedProcess([], 1, "Generating... 'downloadRecordingInfo'", "")
    with pytest.raises(NeedsHuman):
        download._check_output(res)


def test_cookie_help_depends_on_source():
    res = download.subprocess.CompletedProcess([], 1, "ValueError: cookie MoodleSession not set", "")
    with pytest.raises(NeedsHuman) as e:
        download._check_output(res, "webeep")
    assert "MoodleSession" in e.value.action and "ticket" in e.value.action
    assert "JSESSIONID" in download.cookie_help("archives") and "MoodleSession" not in download.cookie_help("txt")


def test_archive_forma_sets_type_and_topic(settings, course_with_links, monkeypatch):
    c, calls = course_with_links, []
    recs = {"2026-05-18 12-32": ID1, "2026-05-19 15-32": ID2, "2026-05-21 10-31": ID3}
    subjects = {"2026-05-18 12-32": "[Laboratorio] Laboratorio 5 EDP: Poisson 2D",
                "2026-05-19 15-32": "[Lezione] Analisi Galerkin-FEM 1", "2026-05-21 10-31": "[Lezione] Analisi Galerkin-FEM 2"}
    monkeypatch.setattr(download, "_run_prd", _fake_prd(recs, calls, subjects))
    rep = download.run(StepContext(settings, c, quiet=True))
    assert sorted(rep.done) == ["2026-05-18_prova_lab01", "2026-05-19_prova_lez01", "2026-05-21_prova_lez02"]
    m = c.cartella / ".sbob" / "manifest.json"
    import json
    assert json.loads(m.read_text())["meta"]["2026-05-19_prova_lez01"]["argomento"] == "Analisi Galerkin-FEM 1"


def test_archives_cookie_help():
    h = download.cookie_help("archives")
    assert "JSESSIONID" in h and "INGRESSCOOKIE" in h and "onlineservices.polimi.it" in h


def test_enriched_links_file_sets_type_topic_and_date(settings, course_with_links, monkeypatch):
    c, calls = course_with_links, []
    (c.cartella / "link.txt").write_text(
        "# commento\n"
        f"https://x.webex.com/recordingservice/sites/x/recording/{ID1}/playback\t08/05/2026 10:33\tLezione\tDF Trasporto 4\n"
        f"https://x.webex.com/recordingservice/sites/x/recording/{ID2}/playback\t07/05/2026 10:31\tLaboratorio\tLab 3\n")
    # Webex dà date sbagliate (10/05 01:45 e 16/05): devono vincere quelle dell'archivio
    recs = {"2026-05-10 01-45": ID1, "2026-05-16 17-20": ID2}
    monkeypatch.setattr(download, "_run_prd", _fake_prd(recs, calls))
    rep = download.run(StepContext(settings, c, quiet=True))
    assert sorted(rep.done) == ["2026-05-07_prova_lab01", "2026-05-08_prova_lez01"]
    sent = open(calls[0][1]).read().split()
    assert len(sent) == 2 and all("\t" not in l and not l.startswith("#") for l in sent)   # a prd solo i link


def test_prd_command_modes(settings, tmp_path, monkeypatch):
    clone = tmp_path / "clone"
    (clone / ".venv" / "bin").mkdir(parents=True)
    (clone / ".venv" / "bin" / "python").write_text("")
    settings.downloader = str(clone)
    cmd, env = download.prd_command(settings)
    assert cmd[0].endswith(".venv/bin/python") and str(clone) in env["PYTHONPATH"]       # venv locale

    monkeypatch.setattr(download.shutil, "which", lambda n: "/usr/bin/uv")
    settings.downloader = "git+https://github.com/x/y@b"
    cmd, _ = download.prd_command(settings)
    assert cmd[:5] == ["/usr/bin/uv", "run", "--no-project", "--quiet", "--with"] and cmd[5] == "git+https://github.com/x/y@b"

    settings.downloader = str(tmp_path / "manca")
    with pytest.raises(NeedsHuman):
        download.prd_command(settings)
    monkeypatch.setattr(download.shutil, "which", lambda n: None)
    settings.downloader = "git+https://github.com/x/y"
    with pytest.raises(NeedsHuman) as e:
        download.prd_command(settings)
    assert "uv" in e.value.action


def test_expired_cookie_triggers_silent_renew_and_retry(settings, course_with_links, monkeypatch):
    c, calls = course_with_links, []
    good = _fake_prd({"2025-09-17 09-30": ID1}, calls)
    state = {"n": 0}

    def flaky(settings, args, cwd):
        state["n"] += 1
        if state["n"] == 1:                                   # primo tentativo: ticket scaduto
            return download.subprocess.CompletedProcess(args, 1, "Try refreshing the ticket", "")
        return good(settings, args, cwd)
    monkeypatch.setattr(download, "_run_prd", flaky)
    monkeypatch.setattr(download, "try_renew_login", lambda ctx: True)
    rep = download.run(StepContext(settings, c, dry_run=True, quiet=True))
    assert rep.done == ["2025-09-17_prova_lez01"] and state["n"] == 2

    state["n"] = 0
    monkeypatch.setattr(download, "try_renew_login", lambda ctx: False)  # rinnovo impossibile → serve l'utente
    with pytest.raises(NeedsHuman) as e:
        download.run(StepContext(settings, c, dry_run=True, quiet=True))
    assert e.value.action.startswith("sbob login")
